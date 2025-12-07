from docx import Document
from io import BytesIO
from .style_utils import replace_text_keep_style


def merge_docs_in_memory(file_a_bytes: bytes, file_b_bytes: bytes) -> BytesIO:
    doc_a = Document(BytesIO(file_a_bytes))
    doc_b = Document(BytesIO(file_b_bytes))

    # Loop through tables by index
    for i, table_a in enumerate(doc_a.tables):
        if i >= len(doc_b.tables):
            break

        table_b = doc_b.tables[i]
        rows_a = table_a.rows
        rows_b = table_b.rows

        if len(rows_a) != len(rows_b):
            continue

        # Check header for "translation" keyword
        header_text = " ".join(cell.text for cell in rows_a[0].cells)
        if "translation" not in header_text.lower():
            continue

        # Skip header row
        for row_a, row_b in zip(rows_a[1:], rows_b[1:]):
            src = row_b.cells[-1].text
            replace_text_keep_style(row_a.cells[-1], src)

    # Output to memory
    output = BytesIO()
    doc_a.save(output)
    output.seek(0)
    return output
